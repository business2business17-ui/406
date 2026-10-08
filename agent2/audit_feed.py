"""Reopen-and-verify audit of a generated feed (Agent 2, sections 32/118-120). Usage: audit_feed.py <part>"""
import collections, datetime, hashlib, json, os, sys, warnings, zipfile
import openpyxl
warnings.simplefilter("ignore")
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import feed_mapping as fm
ROOT = os.path.dirname(HERE); G = os.path.join(HERE, "feeds", "generated") + "/"
part = sys.argv[1]; BATCH = "US-AIRPODS-20261007-001"
name = f"AmazonFeed_US_PORTABLE_ELECTRONIC_DEVICE_COVER_{BATCH}_{part}.xlsm"
tpl = os.path.join(HERE, "templates", "raw", "PORTABLE_ELECTRONIC_DEVICE_COVER_2026-10-08.xlsm"); f = G + name
sha = lambda p: hashlib.sha256(open(p, "rb").read()).hexdigest()
a, b = zipfile.ZipFile(tpl), zipfile.ZipFile(f)
vba = lambda z: sorted(n for n in z.namelist() if "vba" in n.lower())
c = {"zip_ok": b.testzip() is None and a.namelist() == b.namelist(),
     "non_template_parts_identical": all(a.read(n) == b.read(n) for n in a.namelist() if n != "xl/worksheets/sheet5.xml"),
     "macro_parts_identical_to_template": vba(a) == vba(b) and a.read("[Content_Types].xml") == b.read("[Content_Types].xml")}
A, B = openpyxl.load_workbook(tpl, keep_vba=True), openpyxl.load_workbook(f, keep_vba=True)
c["sheet_names_and_states_same"] = A.sheetnames == B.sheetnames and [s.sheet_state for s in A] == [s.sheet_state for s in B]
At, Bt = A["Template"], B["Template"]
c["rows1_6_identical"] = all(At.cell(r, k).value == Bt.cell(r, k).value for r in range(1, 7) for k in range(1, 317))
recs = [json.loads(l) for l in open(os.path.join(ROOT, "output/jsonl/handoff_US_batch001.jsonl"), encoding="utf-8")]
cn = openpyxl.utils.column_index_from_string
res = json.load(open(G + "compile_result.json"))
mism = sum(1 for i, x in enumerate(recs) if Bt.cell(7 + i, 1).value != x["sku"] or Bt.cell(7 + i, 7).value != x["content"]["title"])
minfail = [(r, L) for r in range(7, 137) for L in ("A", "B", "C", "G", "I", "J") if Bt.cell(r, cn(L)).value in (None, "")]
minfail += [(r, "J!=GTIN Exempt") for r in range(7, 137) if Bt.cell(r, cn("J")).value != "GTIN Exempt"]
c.update(row_alignment_mismatch=mism, minimum_field_failures=len(minfail),
         sku_unique=len({Bt.cell(r, 1).value for r in range(7, 137)}) == 130, rows_written=sum(1 for r in range(7, 200) if Bt.cell(r, 1).value),
         formula_cells_in_data=sum(1 for r in range(7, 137) for k in range(1, 317) if str(Bt.cell(r, k).value or "").startswith("=")),
         text_roundtrip_bad=sum(1 for i, x in enumerate(recs) for L, v in (("H", x["content"]["item_highlights"]), ("AD", x["content"]["description"]), ("AJ", x["content"]["backend_search_terms"])) if Bt.cell(7 + i, cn(L)).value != v),
         template_rule_and_enum_issues=len(res["issues"]),
         pro_photo5_present=sum(1 for r in range(7, 72) for k in range(20, 28) if str(Bt.cell(r, k).value or "").endswith("_5.jpg")),
         airpods_pro_3rd_used=sum(1 for r in range(7, 137) for k in range(cn("CP"), cn("CT") + 1) if "3rd" in str(Bt.cell(r, k).value or "")))
pr = collections.Counter(tuple(Bt.cell(r, cn(L)).value for L in ("EV", "EZ", "FE", "EX", "EY", "FL", "FN", "FF", "FG")) for r in range(7, 137))
c["price_tuples(EV,EZ,FE,EX,EY,FL,FN,FF,FG)"] = [list(k) + [v] for k, v in pr.items()]
(ev, ez, fe, ex, ey, fl, fn, ff, fg), = [k for k in pr][:1]
c["price_logic"] = bool(ex <= ez <= ev <= ey and fn < fl < fe <= ez and ff is None and fg is None)
ok = (c["zip_ok"] and c["non_template_parts_identical"] and c["macro_parts_identical_to_template"] and c["sheet_names_and_states_same"] and c["rows1_6_identical"]
      and mism == 0 and not minfail and c["sku_unique"] and c["rows_written"] == 130 and c["formula_cells_in_data"] == 0 and c["text_roundtrip_bad"] == 0
      and c["template_rule_and_enum_issues"] == 0 and c["pro_photo5_present"] == 0 and c["airpods_pro_3rd_used"] == 0 and c["price_logic"] and len(pr) == 1)
c["final_audit"] = "PASS" if ok else "FAIL"; status = "READY_FOR_AMAZON_UPLOAD" if ok else "NOT_READY_FOR_AMAZON_UPLOAD"
print(json.dumps(c, ensure_ascii=False, indent=1))
man = {"feed_filename": name, "file_status": status, "marketplace": "US", "product_type": "PORTABLE_ELECTRONIC_DEVICE_COVER", "batch_id": BATCH, "part_number": int(part), "sku_count": 130,
       "operation_counts": {"CREATE": 130}, "blocked_count": 0, "generated_at": datetime.datetime.utcnow().isoformat() + "Z", "agent2_version": "1.0.0", "handoff_schema_version": "1.0.0",
       "agent1_record_version": recs[0]["record_version"], "pricing_policy_version": "2026-10-08-v3 with user overrides (tiers from Business Price, min = 4-pc tier price)",
       "template": {"path": os.path.relpath(tpl, ROOT), "sha256": sha(tpl)}, "identifier_mode": "GTIN_EXEMPT", "execution_mode": "HYBRID (USER_UPLOAD template + GITHUB branch claude/clever-albattani-5szkv4)",
       "feed_sha256": sha(f), "agent1_jsonl_sha256": sha(os.path.join(ROOT, "output/jsonl/handoff_US_batch001.jsonl")), "audit": c}
json.dump(man, open(G + f"Feed_Manifest_{BATCH}.json", "w"), ensure_ascii=False, indent=1)
json.dump({"file_status": status, "audit": c, "issues": res["issues"]}, open(G + f"Feed_Validation_{BATCH}.json", "w"), ensure_ascii=False, indent=1)
open(G + f"Cell_Provenance_{BATCH}.jsonl", "w", encoding="utf-8").write("\n".join(json.dumps(p, ensure_ascii=False) for p in res["provenance"]))
open(G + f"Mutation_Log_{BATCH}.jsonl", "w", encoding="utf-8").write("\n".join(json.dumps(p, ensure_ascii=False) for p in res["mutations"]))
json.dump({"status": status, "feed": name, "feed_sha256": sha(f), "template_sha256": sha(tpl), "rows": 130, "identifier_mode": "GTIN Exempt"}, open(G + "Feed_Status.json", "w"), ensure_ascii=False, indent=1)
print(status)
